"""读取飞书知识库里的电子表格（以应用身份调用开放平台 API）。"""
import os
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import requests

API = "https://open.feishu.cn/open-apis"
CONFIG_KEYS = ("FEISHU_APP_ID", "FEISHU_APP_SECRET", "WIKI_TOKEN", "SHEET_ID", "DASHBOARD_URL", "LOCAL_XLSX", "PUBLIC_URL")


def load_config(path=Path(__file__).with_name("config.env")):
    """读取 config.env（KEY=VALUE 格式），环境变量优先（云端运行时只有环境变量）。"""
    cfg = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                cfg[k.strip()] = v.strip()
    for k in {*cfg, *CONFIG_KEYS}:
        if os.environ.get(k):
            cfg[k] = os.environ[k]
    return cfg


class FeishuSheet:
    def __init__(self, app_id, app_secret):
        self.app_id = app_id
        self.app_secret = app_secret
        self._tok = (None, 0)
        self._lock = threading.Lock()

    def _token(self):
        """tenant_access_token 有效期 2 小时，缓存复用，提前 5 分钟换新。"""
        with self._lock:
            tok, exp = self._tok
            if tok and time.time() < exp:
                return tok
            r = requests.post(f"{API}/auth/v3/tenant_access_token/internal",
                              json={"app_id": self.app_id, "app_secret": self.app_secret}, timeout=10)
            data = r.json()
            if data.get("code") != 0:
                raise RuntimeError(f"获取 tenant_access_token 失败：{data.get('msg')}")
            self._tok = (data["tenant_access_token"], time.time() + data.get("expire", 7200) - 300)
            return self._tok[0]

    def _get(self, path, **params):
        for attempt in range(3):  # 大表偶尔读超时，重试两次
            try:
                r = requests.get(f"{API}{path}", params=params,
                                 headers={"Authorization": f"Bearer {self._token()}"}, timeout=60)
                break
            except requests.exceptions.RequestException:
                if attempt == 2:
                    raise
        data = r.json()
        if data.get("code") != 0:
            raise RuntimeError(f"{path} 调用失败（code={data.get('code')}）：{data.get('msg')}")
        return data["data"]

    def spreadsheet_token(self, wiki_token):
        """知识库节点 token → 电子表格 token。"""
        node = self._get("/wiki/v2/spaces/get_node", token=wiki_token, obj_type="wiki")["node"]
        if node["obj_type"] != "sheet":
            raise RuntimeError(f"该知识库节点不是电子表格，而是 {node['obj_type']}")
        return node["obj_token"]

    def sheets(self, spreadsheet):
        """列出所有工作表：[{sheet_id, title, grid_properties:{row_count, column_count}}]"""
        return self._get(f"/sheets/v3/spreadsheets/{spreadsheet}/sheets/query")["sheets"]

    def values(self, spreadsheet, sheet, max_rows=None, batch_rows=250, workers=8):
        """读取整张工作表的单元格（二维列表），分小批并行读取（公式多的表按大批读很慢）。"""
        grid = sheet["grid_properties"]
        rows = grid["row_count"] if max_rows is None else min(max_rows, grid["row_count"])
        last_col = _col_letter(grid["column_count"])

        def batch(start):
            rng = f"{sheet['sheet_id']}!A{start}:{last_col}{min(start + batch_rows - 1, rows)}"
            data = self._get(f"/sheets/v2/spreadsheets/{spreadsheet}/values/{rng}",
                             valueRenderOption="ToString", dateTimeRenderOption="FormattedString")
            return data["valueRange"].get("values") or []

        with ThreadPoolExecutor(max_workers=workers) as pool:
            out = [row for part in pool.map(batch, range(1, rows + 1, batch_rows)) for row in part]
        # 去掉末尾整行为空的行
        while out and not any(c not in (None, "") for c in out[-1]):
            out.pop()
        return out

    def records(self, spreadsheet, sheet, header_row=1):
        """把工作表读成 [{表头: 值}]，表头取第 header_row 行。"""
        vals = self.values(spreadsheet, sheet)
        if len(vals) < header_row:
            return []
        header = [str(h).strip() if h is not None else "" for h in vals[header_row - 1]]
        return [{h: row[i] if i < len(row) else None for i, h in enumerate(header) if h}
                for row in vals[header_row:]]


    def tables(self, spreadsheet, names):
        """按工作表名读取多张表：{表名: [{表头: 值}]}。"""
        by_title = {s["title"]: s for s in self.sheets(spreadsheet)}
        missing = [n for n in names if n not in by_title]
        if missing:
            raise RuntimeError(f"表格里找不到工作表：{'、'.join(missing)}")
        with ThreadPoolExecutor(max_workers=len(names)) as pool:
            return dict(zip(names, pool.map(lambda n: self.records(spreadsheet, by_title[n]), names)))


def xlsx_tables(path, names):
    """从本地 Excel（飞书导出）读取多张表，格式同 FeishuSheet.tables。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for n in names:
        ws = wb[n]
        ws.reset_dimensions()  # 飞书导出的文件尺寸信息不准，重新扫描
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        out[n] = [{h: r[i] if i < len(r) else None for i, h in enumerate(header) if h}
                  for r in rows if any(c not in (None, "") for c in r)]
    return out


def _col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s
